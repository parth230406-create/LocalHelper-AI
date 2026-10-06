from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

@dataclass
class DocUnit:
    source_name: str
    loc_kind: str
    loc: int
    text: str

def _clean(t: str) -> str:
    return " ".join((t or "").split()).strip()

def extract_units(path: str) -> list[DocUnit]:
    p = Path(path)
    ext = p.suffix.lower()
    try:
        if ext == ".pdf":
            reader = PdfReader(str(p))
            units = []
            for i, page in enumerate(reader.pages, 1):
                t = _clean(page.extract_text() or "")
                if t:
                    units.append(DocUnit(p.name, "page", i, t))
            return units
        if ext in (".txt", ".md"):
            t = _clean(p.read_text(encoding="utf-8", errors="ignore"))
            return [DocUnit(p.name, "text", 1, t)] if t else []
        if ext == ".docx":
            doc = Document(str(p))
            t = _clean("\n".join(para.text for para in doc.paragraphs if para.text.strip()))
            return [DocUnit(p.name, "doc", 1, t)] if t else []
        if ext == ".pptx":
            prs = Presentation(str(p))
            units = []
            for i, slide in enumerate(prs.slides, 1):
                texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
                t = _clean("\n".join(texts))
                if t:
                    units.append(DocUnit(p.name, "slide", i, t))
            return units
    except Exception as e:
        print(f"Error extracting {p.name}: {e}")
    return []

def chunk_text(text: str, size: int = 1500, overlap: int = 100) -> list[str]:
    text = (text or "").strip()
    if not text:
        return []
    if size <= overlap:
        overlap = size // 4
    chunks, start, n = [], 0, len(text)
    while start < n:
        end = min(n, start + size)
        chunks.append(text[start:end])
        if end == n:
            break
        start = end - overlap
    return chunks